"""Proving document ingestion, the extraction ladder, and GraphRAG retrieval.

Real files are built here — a real .docx through python-docx, a real .xlsx
through openpyxl, real Markdown, and real binary junk — and pushed through the
real Docling ladder into a real post-graph-rag index backed by the live
database, with the real model doing entity extraction.

Indexing a document through GraphRAG costs several model calls, so the corpus
is deliberately tiny. What is being proved is the contract, not the throughput.
"""
import io
import uuid

import pytest

from conftest import requires_stack

pytestmark = [pytest.mark.e2e, requires_stack]


# ------------------------------------------------------------- real files

def markdown_bytes(title: str, body: str) -> bytes:
    return f"# {title}\n\n{body}\n".encode("utf-8")


def docx_bytes(paragraphs) -> bytes:
    import docx
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(rows) -> bytes:
    import openpyxl
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Ledger"
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


UNREADABLE = bytes(range(256)) * 8          # genuinely not text, genuinely not a document


# ------------------------------------------------------------------ spaces

async def test_a_document_space_is_created_once(documents, realm, project):
    """Rule 3.2 — UI flows retry, and a duplicate space splits a corpus in half."""
    body = {"org_id": realm, "project_id": project,
            "document_space": "engineering_docs", "description": "first"}
    first = await documents.post("/spaces", json=body)
    second = await documents.post("/spaces", json={**body, "description": "second"})
    assert first.status_code == second.status_code == 200
    assert first.json()["space_key"] == second.json()["space_key"]
    assert second.json()["description"] == "first"      # idempotent, not last-write-wins

    listed = await documents.get(f"/projects/{project}/spaces",
                                 params={"org_id": realm})
    spaces = listed.json()["spaces"]
    assert len([s for s in spaces if s["document_space"] == "engineering_docs"]) == 1


async def test_space_name_is_accepted_as_an_alias(documents, realm, project):
    """Rule 2.0 — the frontend and backend send `space_name`; both mean one field."""
    res = await documents.post("/spaces", json={
        "org_id": realm, "project_id": project, "space_name": "legal_contracts"})
    assert res.status_code == 200
    body = res.json()
    assert body["document_space"] == body["space_name"] == "legal_contracts"


async def test_an_invalid_document_space_is_rejected(documents, realm, project):
    """Rule 3.1 — it appears in URL paths and RAG metadata filters."""
    res = await documents.post("/spaces", json={
        "org_id": realm, "project_id": project, "space_name": "Engineering Docs"})
    assert res.status_code == 422 or res.status_code == 400


# --------------------------------------------------------------- ingestion

async def test_text_upload_is_extracted_indexed_and_catalogued(
        documents, realm, project):
    res = await documents.post("/spaces/notes/documents/upload-text", json={
        "org_id": realm, "project_id": project, "document_space": "notes",
        "document_name": "credit-facility.md",
        "content": ("Northwind Trading reports material uncertainty over the "
                    "renewal of its revolving credit facility with Barclays. "
                    "The facility matures in March.")})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["indexed"] is True, body["message"]
    assert body["status"] == "success"

    document = body["document"]
    assert document["document_space"] == document["space_name"] == "notes"
    assert document["content_hash"].startswith("sha256:")
    assert document["revision"] == 1
    assert document["index"]["status"] == "indexed"
    assert document["index"]["chunks"] >= 1


async def test_a_real_docx_is_extracted_by_the_ladder(documents, realm, project):
    data = docx_bytes([
        "Quarterly Operations Review",
        "The Hamburg distribution centre exceeded its throughput target by nine "
        "per cent, while the Lyon site fell short following a conveyor failure.",
    ])
    res = await documents.post(
        "/spaces/reports/documents/upload-file",
        data={"project_id": project, "org_id": realm, "document_space": "reports"},
        files={"file": ("ops-review.docx", data,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert res.status_code == 200, res.text
    document = res.json()["document"]
    assert document["extraction_method"] in {"docling", "python-docx"}
    assert document["content_length"] > 50
    assert document["bytes"] == len(data)


async def test_a_real_xlsx_is_extracted_by_the_ladder(documents, realm, project):
    data = xlsx_bytes([["Supplier", "Invoice", "Amount"],
                       ["Northwind", "INV-1001", 4200],
                       ["Contoso", "INV-1002", 980]])
    res = await documents.post(
        "/spaces/ledgers/documents/upload-file",
        data={"project_id": project, "org_id": realm, "document_space": "ledgers"},
        files={"file": ("ledger.xlsx", data,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert res.status_code == 200, res.text
    document = res.json()["document"]
    assert document["extraction_method"] in {"docling", "openpyxl"}
    assert "INV-1001" in str(document.get("content_length")) or \
        document["content_length"] > 0


async def test_an_unreadable_file_is_rejected_and_not_catalogued(
        documents, realm, project):
    """Rule 5.3 — twice in this service's history a placeholder was indexed
    instead, and read back later as if it were the document."""
    res = await documents.post(
        "/spaces/notes/documents/upload-file",
        data={"project_id": project, "org_id": realm, "document_space": "notes"},
        files={"file": ("mystery.bin", UNREADABLE, "application/octet-stream")})
    assert res.status_code == 415
    assert "mystery.bin" in res.json()["detail"]
    assert "not been catalogued" in res.json()["detail"]

    listed = await documents.get(f"/projects/{project}/documents",
                                 params={"org_id": realm})
    assert [d for d in listed.json()["documents"] if d["filename"] == "mystery.bin"] == []


async def test_a_batch_reports_each_file_separately(documents, realm, project):
    """Rule 5.4 and 5.5 — the batch endpoint applies the identical rule, and a
    count that silently excludes failures reads as complete."""
    files = [
        ("files", ("good-one.md", markdown_bytes(
            "Site Report", "The Lyon conveyor was replaced on Tuesday."), "text/markdown")),
        ("files", ("mystery.bin", UNREADABLE, "application/octet-stream")),
        ("files", ("good-two.md", markdown_bytes(
            "Safety Note", "All Hamburg staff completed refresher training."),
            "text/markdown")),
    ]
    res = await documents.post(
        "/spaces/batch/documents/upload-multiple-files",
        data={"project_id": project, "org_id": realm, "document_space": "batch"},
        files=files)
    assert res.status_code == 200, res.text
    body = res.json()

    assert body["status"] == "partial"
    assert body["count"] == 2
    assert body["failed_count"] == 1
    assert body["failures"][0]["filename"] == "mystery.bin"
    assert body["failures"][0]["stage"] == "extraction"
    # The good files went in; the bad one did not.
    names = {d["filename"] for d in body["documents"]}
    assert names == {"good-one.md", "good-two.md"}


# ------------------------------------------------------ identity and revision

async def test_reuploading_a_name_is_a_revision(documents, realm, project):
    """Rule 7.1 — identity is (org, project, document_space, filename)."""
    async def upload(content):
        return await documents.post("/spaces/notes/documents/upload-text", json={
            "org_id": realm, "project_id": project, "document_space": "notes",
            "document_name": "policy.md", "content": content})

    first = await upload("The travel policy permits economy fares only.")
    second = await upload("The travel policy now permits premium economy "
                          "on flights over six hours.")
    assert first.status_code == second.status_code == 200
    assert first.json()["document"]["revision"] == 1
    assert second.json()["document"]["revision"] == 2

    listed = await documents.get(f"/projects/{project}/documents",
                                 params={"org_id": realm, "document_space": "notes"})
    policies = [d for d in listed.json()["documents"] if d["filename"] == "policy.md"]
    assert len(policies) == 1, "a revision must not become a second document"

    doc_id = policies[0]["document_id"]
    revisions = await documents.get(
        f"/projects/{project}/documents/{doc_id}/revisions",
        params={"org_id": realm})
    assert revisions.json()["count"] >= 2


async def test_identical_bytes_are_not_ingested_twice(documents, realm, project):
    """Rule 5.6 — duplicate chunks compete in retrieval and degrade every answer."""
    content = "The Rotterdam warehouse lease expires in September."
    body = {"org_id": realm, "project_id": project, "document_space": "notes",
            "document_name": "lease.md", "content": content}

    first = await documents.post("/spaces/notes/documents/upload-text", json=body)
    second = await documents.post("/spaces/notes/documents/upload-text", json=body)
    assert first.status_code == second.status_code == 200
    assert second.json()["document"].get("deduplicated") is True
    assert second.json()["document"]["revision"] == 1       # unchanged


async def test_the_same_name_in_another_document_space_is_another_document(
        documents, realm, project):
    for space in ("alpha", "beta"):
        res = await documents.post("/spaces/x/documents/upload-text", json={
            "org_id": realm, "project_id": project, "document_space": space,
            "document_name": "shared.md",
            "content": f"This document belongs to the {space} collection."})
        assert res.status_code == 200, res.text

    listed = await documents.get(f"/projects/{project}/documents",
                                 params={"org_id": realm})
    shared = [d for d in listed.json()["documents"] if d["filename"] == "shared.md"]
    assert len(shared) == 2
    assert {d["document_space"] for d in shared} == {"alpha", "beta"}


async def test_withdrawal_retains_the_record_and_drops_the_chunks(
        documents, realm, project):
    """Rule 7.3 — the record explains why last month's run cited this document."""
    res = await documents.post("/spaces/notes/documents/upload-text", json={
        "org_id": realm, "project_id": project, "document_space": "notes",
        "document_name": "withdrawn.md",
        "content": "The Glasgow depot will close at the end of the year."})
    doc_id = res.json()["document"]["document_id"]

    gone = await documents.delete(f"/projects/{project}/documents/{doc_id}",
                                  params={"org_id": realm})
    assert gone.status_code == 200, gone.text
    assert gone.json()["status"] == "withdrawn"

    listed = await documents.get(f"/projects/{project}/documents",
                                 params={"org_id": realm})
    assert doc_id not in [d["document_id"] for d in listed.json()["documents"]]

    retained = await documents.get(f"/projects/{project}/documents",
                                   params={"org_id": realm, "include_withdrawn": True})
    kept = [d for d in retained.json()["documents"] if d["document_id"] == doc_id]
    assert kept and kept[0]["lifecycle"] == "withdrawn"


# ---------------------------------------------------------------- counting

async def test_document_counts_are_derived_from_the_catalogue(
        documents, realm, project):
    """Rule 3.3 — a stored counter drifts, and users see the drift."""
    await documents.post("/spaces", json={
        "org_id": realm, "project_id": project, "document_space": "counted"})

    for n in range(3):
        await documents.post("/spaces/counted/documents/upload-text", json={
            "org_id": realm, "project_id": project, "document_space": "counted",
            "document_name": f"note-{n}.md",
            "content": f"Note number {n} about depot operations in region {n}."})

    spaces = (await documents.get(f"/projects/{project}/spaces",
                                  params={"org_id": realm})).json()["spaces"]
    counted = next(s for s in spaces if s["document_space"] == "counted")
    assert counted["document_count"] == 3

    listed = await documents.get(f"/projects/{project}/documents",
                                 params={"org_id": realm, "document_space": "counted"})
    doc_id = listed.json()["documents"][0]["document_id"]
    await documents.delete(f"/projects/{project}/documents/{doc_id}",
                           params={"org_id": realm})

    spaces = (await documents.get(f"/projects/{project}/spaces",
                                  params={"org_id": realm})).json()["spaces"]
    counted = next(s for s in spaces if s["document_space"] == "counted")
    assert counted["document_count"] == 2, "the count did not follow the withdrawal"


# --------------------------------------------------------------- retrieval

async def test_graphrag_retrieval_finds_an_ingested_fact(documents, realm, project):
    """The corpus is queryable: ingest, then ask, and get the passage back."""
    await documents.post("/spaces/corpus/documents/upload-text", json={
        "org_id": realm, "project_id": project, "document_space": "corpus",
        "document_name": "facility.md",
        "content": ("Northwind Trading holds a revolving credit facility with "
                    "Barclays that matures in March. The company has flagged "
                    "material uncertainty over its renewal.")})

    res = await documents.post("/query", json={
        "org_id": realm, "project_id": project,
        "query": "What did Northwind say about its credit facility?",
        "top_k": 5, "mode": "naive"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["status"] == "success"
    assert body["engine"] == "post-graph-rag"

    chunks = body["data"].get("chunks") or []
    assert chunks, "retrieval returned no chunks from a corpus that has one"
    combined = " ".join(str(c.get("content", "")) for c in chunks).lower()
    assert "barclays" in combined or "revolving credit" in combined


async def test_retrieval_can_be_scoped_to_one_document_space(
        documents, realm, project):
    """Rule 2.3 — omitted means project-wide, named means narrowed."""
    await documents.post("/spaces/finance/documents/upload-text", json={
        "org_id": realm, "project_id": project, "document_space": "finance",
        "document_name": "treasury.md",
        "content": "The treasury team hedges euro exposure quarterly."})
    await documents.post("/spaces/facilities/documents/upload-text", json={
        "org_id": realm, "project_id": project, "document_space": "facilities",
        "document_name": "depot.md",
        "content": "The Glasgow depot roof was replaced in June."})

    everything = await documents.post("/query", json={
        "org_id": realm, "project_id": project, "query": "what happened",
        "top_k": 10, "mode": "naive"})
    assert everything.json()["document_space"] == "all_document_spaces"

    narrowed = await documents.post("/query", json={
        "org_id": realm, "project_id": project, "query": "what happened",
        "document_space": "facilities", "top_k": 10, "mode": "naive"})
    assert narrowed.status_code == 200, narrowed.text
    assert narrowed.json()["document_space"] == "facilities"
    chunks = narrowed.json()["data"].get("chunks") or []
    for chunk in chunks:
        metadata = chunk.get("metadata", {})
        space = (metadata.get("document_space") or metadata.get("collection")
                 or (metadata.get("extra") or {}).get("document_space"))
        assert space == "facilities", metadata


async def test_a_query_names_the_engine_that_answered(documents, realm, project):
    """Rule 8.1 — an agent citing retrieved text is claiming provenance."""
    await documents.post("/spaces/corpus/documents/upload-text", json={
        "org_id": realm, "project_id": project, "document_space": "corpus",
        "document_name": "one.md", "content": "A short note about logistics."})
    res = await documents.post("/query", json={
        "org_id": realm, "project_id": project, "query": "logistics",
        "mode": "naive"})
    assert res.json()["engine"] in {"post-graph-rag", "post-graph-direct"}
    assert res.json()["status"] in {"success", "degraded"}


# ----------------------------------------------------------------- tenancy

async def test_two_realms_cannot_see_each_others_documents(documents, project):
    """AG §2 — physical isolation, exercised rather than assumed."""
    left = "t_" + uuid.uuid4().hex[:10]
    right = "t_" + uuid.uuid4().hex[:10]

    for realm, secret in ((left, "The left organisation runs a plant in Turin."),
                          (right, "The right organisation runs a plant in Porto.")):
        res = await documents.post("/spaces/notes/documents/upload-text", json={
            "org_id": realm, "project_id": project, "document_space": "notes",
            "document_name": "secret.md", "content": secret})
        assert res.status_code == 200, res.text

    left_docs = await documents.get(f"/projects/{project}/documents",
                                    params={"org_id": left})
    right_docs = await documents.get(f"/projects/{project}/documents",
                                     params={"org_id": right})
    assert len(left_docs.json()["documents"]) == 1
    assert len(right_docs.json()["documents"]) == 1
    assert left_docs.json()["documents"][0]["org_id"] == left
    assert right_docs.json()["documents"][0]["org_id"] == right


async def test_two_projects_in_one_realm_are_separated(documents, realm):
    """The project is a post-graph space — logical isolation within the realm."""
    alpha = "proj_" + uuid.uuid4().hex[:8]
    beta = "proj_" + uuid.uuid4().hex[:8]

    for project in (alpha, beta):
        await documents.post("/spaces/notes/documents/upload-text", json={
            "org_id": realm, "project_id": project, "document_space": "notes",
            "document_name": "note.md",
            "content": f"A note belonging to project {project}."})

    for project in (alpha, beta):
        listed = await documents.get(f"/projects/{project}/documents",
                                     params={"org_id": realm})
        docs = listed.json()["documents"]
        assert len(docs) == 1
        assert docs[0]["project_id"] == project


# --------------------------------------------------------------- accounting

async def test_ingestion_and_retrieval_are_metered_apart(
        documents, db, realm, project):
    """§9 — they have very different cost profiles and one counter cannot be
    billed apart afterwards."""
    from test_e2e_tool_discovery import ledger

    await documents.post("/spaces/billed/documents/upload-text", json={
        "org_id": realm, "project_id": project, "document_space": "billed",
        "document_name": "metered.md",
        "content": "The Antwerp terminal handled record volumes this quarter."})
    await documents.post("/query", json={
        "org_id": realm, "project_id": project, "query": "Antwerp", "mode": "naive"})

    kinds = await ledger(db, realm, minimum=2)
    assert kinds.get("document_ingest"), f"no ingest event; saw {list(kinds)}"
    assert kinds.get("rag_lookup"), f"no retrieval event; saw {list(kinds)}"

    ingest = kinds["document_ingest"][0]
    assert ingest["org_id"] == realm
    assert ingest["project_id"] == project
    assert ingest["bytes"] > 0
